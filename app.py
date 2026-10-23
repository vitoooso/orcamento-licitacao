"""
app.py — interface do técnico orçamentista (Streamlit).

Três telas:
  Preencher planilha -> envia a planilha da licitação e recebe ela preenchida + relatório
  Cadastrar preços   -> alimenta o banco com planilhas de licitações passadas
  Consultar banco    -> busca preços, exporta e desfaz importações

Regras de negócio: core.py. Banco de dados: storage.py.
Logo: salve o arquivo em assets/logo.png (também aceita .svg, .jpg, .webp).
"""
from __future__ import annotations

import base64
from datetime import date, datetime, timedelta
import html
import io
import os
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import core
import storage

# ----------------------------------------------------------------------------
# Identidade visual
# ----------------------------------------------------------------------------
AZUL_MARINHO = "#0E2A47"
AZUL = "#1565A8"
NEVOA = "#EEF3F8"
TINTA = "#16263A"
CINZA_TEXTO = "#5A6B7D"
BORDA = "#D5DDE6"

# As cores dos status são as MESMAS que o core.py pinta no Excel.
COR_BARRA = {core.ST_OK: AZUL, **core.COR_STATUS}
COR_CELULA_STATUS = {core.ST_OK: "#D6E6F5", **core.COR_STATUS}
COR_SITUACAO = {storage.SIT_NOVO: "#D8EFE0", storage.SIT_EXISTE: "#E3E3E3", storage.SIT_REPETIDO: "#FFE7A8"}
ORDEM_STATUS = [core.ST_OK, core.ST_NAO, core.ST_AMB, core.ST_MESCLADA, core.ST_JA]
LEGENDA_STATUS = {
    core.ST_OK: "preenchidos",
    core.ST_NAO: "sem preço encontrado",
    core.ST_AMB: "com preços conflitantes",
    core.ST_MESCLADA: "com célula de valor mesclada",
    core.ST_JA: "já tinham valor (mantidos)",
}

ONDAS = (
    "url(\"data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' width='160' height='20' viewBox='0 0 160 20'%3E"
    "%3Cpath d='M0 10 Q40 0 80 10 T160 10' fill='none' stroke='%23ffffff' stroke-width='1.4'/%3E%3C/svg%3E\")"
)

CSS = f"""
<style>
/* ---- Estrutura da página ---- */
header[data-testid="stHeader"] {{ background: rgba(255, 255, 255, .9); backdrop-filter: blur(6px); }}
.block-container {{ padding-top: 4.5rem; padding-bottom: 3rem; max-width: 1120px; }}

/* ---- Faixa de topo ---- */
.atl-hero {{
  position: relative; overflow: hidden; display: flex; align-items: center; gap: 1.5rem; flex-wrap: wrap;
  background: {AZUL_MARINHO}; color: #FFFFFF; border-radius: 12px; padding: 1.4rem 1.75rem 2.1rem;
  margin-bottom: .75rem;
}}
.atl-hero::after {{
  content: ""; position: absolute; left: 0; right: 0; bottom: 6px; height: 30px; opacity: .16; pointer-events: none;
  background-image: {ONDAS}, {ONDAS};
  background-repeat: repeat-x; background-size: 160px 20px; background-position: 0 0, 80px 10px;
}}
.atl-logo {{ background: #FFFFFF; border-radius: 8px; padding: .55rem .85rem; display: flex; align-items: center; }}
.atl-logo img {{ height: 46px; width: auto; display: block; }}
.atl-marca {{ color: {AZUL_MARINHO}; font-weight: 700; font-size: 1.05rem; line-height: 1.15; }}
.atl-marca small {{ display: block; font-weight: 500; font-size: .78rem; color: {CINZA_TEXTO}; }}
.atl-hero-texto .nome {{ font-size: 1.55rem; font-weight: 700; line-height: 1.2; letter-spacing: -.01em; }}
.atl-hero-texto .sub {{ font-size: .95rem; opacity: .8; margin-top: .15rem; }}
.atl-banco {{
  margin-left: auto; display: inline-flex; align-items: center; gap: .5rem; font-size: .85rem;
  padding: .4rem .85rem; border-radius: 999px; background: rgba(255, 255, 255, .1);
  border: 1px solid rgba(255, 255, 255, .22); white-space: nowrap;
}}
.atl-banco::before {{ content: ""; width: .55rem; height: .55rem; border-radius: 50%; background: var(--cor-ponto); }}

/* ---- Abas ---- */
.stTabs [data-baseweb="tab-list"] {{ gap: .5rem; border-bottom: 1px solid {BORDA}; }}
.stTabs [data-baseweb="tab"] {{ padding: .85rem .9rem .7rem; }}
.stTabs [data-baseweb="tab"] p {{ font-size: 1rem; font-weight: 500; }}
.stTabs [aria-selected="true"] p {{ color: {AZUL_MARINHO}; font-weight: 600; }}
.stTabs [data-baseweb="tab-panel"] {{ padding-top: 1.75rem; }}

/* ---- Passos em linha do tempo ---- */
[class*="st-key-etapa_"] {{ position: relative; padding: 0 0 2rem 3.25rem; }}
[class*="st-key-etapa_"]::before {{
  content: ""; position: absolute; left: 1.05rem; top: 2.6rem; bottom: .4rem; width: 2px; background: {BORDA};
}}
.atl-etapa-cab {{ display: flex; align-items: center; gap: 1rem; margin: 0 0 .35rem -3.25rem; }}
.atl-num {{
  flex: none; width: 2.2rem; height: 2.2rem; border-radius: 50%; display: grid; place-items: center;
  background: {AZUL_MARINHO}; color: #FFFFFF; font-weight: 600; font-size: 1rem;
}}
.atl-etapa-titulo {{ font-size: 1.22rem; font-weight: 600; color: {TINTA}; line-height: 1.25; }}
.atl-etapa-desc {{ font-size: .92rem; color: {CINZA_TEXTO}; margin-top: .1rem; max-width: 70ch; }}

/* ---- Envio de arquivo e botões ---- */
[data-testid="stFileUploaderDropzone"] {{
  border: 1.5px dashed #A9BBCD; background: #F6F9FC; border-radius: 10px; padding: 1.1rem 1.25rem;
}}
[data-testid="stBaseButton-primary"] {{ padding: .6rem 1.4rem; font-weight: 600; }}

/* ---- Como funciona (tela inicial) ---- */
.atl-como {{ display: grid; grid-template-columns: repeat(3, 1fr); gap: 1rem 2rem; margin: 1.1rem 0 .25rem; }}
.atl-como div {{ border-left: 3px solid {AZUL}; padding: .1rem 0 .1rem .9rem; }}
.atl-como b {{ display: block; color: {TINTA}; font-weight: 600; margin-bottom: .15rem; }}
.atl-como span {{ color: {CINZA_TEXTO}; font-size: .92rem; line-height: 1.45; }}

/* ---- Resultado ---- */
.atl-resultado {{ background: {NEVOA}; border-radius: 10px; padding: 1.1rem 1.35rem .9rem; margin-bottom: 1rem; }}
.atl-resumo {{ font-size: 1.05rem; color: {TINTA}; margin: 0; }}
.atl-resumo strong {{ font-size: 2.1rem; color: {AZUL_MARINHO}; font-weight: 700; margin-right: .4rem; }}
.atl-barra {{
  display: flex; height: 22px; border-radius: 4px; overflow: hidden; background: #DDE4EB;
  margin: .65rem 0 .75rem; border: 1px solid #CBD5E0;
}}
.atl-barra span {{ display: block; height: 100%; }}
.atl-legenda {{ display: flex; flex-wrap: wrap; gap: .35rem 1.4rem; font-size: .9rem; color: {TINTA}; }}
.atl-legenda i {{
  display: inline-block; width: .8rem; height: .8rem; border-radius: 2px; margin-right: .4rem;
  vertical-align: -1px; border: 1px solid rgba(0, 0, 0, .18);
}}

/* ---- Estados vazios e rodapé ---- */
.atl-vazio {{
  border: 1px solid {BORDA}; border-radius: 10px; padding: 2rem 2rem 1.75rem; max-width: 640px;
  background: #FFFFFF;
}}
.atl-vazio .titulo {{ font-size: 1.15rem; font-weight: 600; color: {TINTA}; margin-bottom: .35rem; }}
.atl-vazio p {{ color: {CINZA_TEXTO}; margin: 0; line-height: 1.55; }}
.atl-rodape {{
  margin-top: 3rem; padding-top: 1rem; border-top: 1px solid {BORDA}; font-size: .82rem; color: {CINZA_TEXTO};
}}

@media (max-width: 720px) {{
  .atl-hero {{ padding: 1.2rem 1.2rem 2rem; }}
  .atl-banco {{ margin-left: 0; }}
  .atl-como {{ grid-template-columns: 1fr; }}
  [class*="st-key-etapa_"] {{ padding-left: 2.75rem; }}
  .atl-etapa-cab {{ margin-left: -2.75rem; }}
}}
</style>
"""


# ----------------------------------------------------------------------------
# Formatação
# ----------------------------------------------------------------------------
def fmt_int(n: int) -> str:
    return f"{n:,}".replace(",", ".")


def brl(valor) -> str:
    if valor is None or pd.isna(valor):
        return ""
    return "R$ " + f"{float(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def letra(col: int | None) -> str:
    return "" if col is None else get_column_letter(col + 1)


def arquivo_logo() -> Path | None:
    pasta = Path(__file__).parent / "assets"
    for ext in ("png", "svg", "jpg", "jpeg", "webp"):
        caminho = pasta / f"logo.{ext}"
        if caminho.exists():
            return caminho
    return None


def logo_html() -> str:
    caminho = arquivo_logo()
    if caminho is None:
        return '<span class="atl-marca">Atlântico<small>Engenharia e Inovação</small></span>'
    mime = {"svg": "image/svg+xml", "jpg": "image/jpeg", "jpeg": "image/jpeg"}.get(caminho.suffix[1:], f"image/{caminho.suffix[1:]}")
    b64 = base64.b64encode(caminho.read_bytes()).decode()
    return f'<img src="data:{mime};base64,{b64}" alt="Atlântico Engenharia">'


def avisar(chave: str, mensagem: str) -> None:
    """Guarda uma mensagem de sucesso para mostrar depois do st.rerun()."""
    st.session_state[f"aviso_{chave}"] = mensagem


def mostrar_aviso(chave: str) -> None:
    mensagem = st.session_state.pop(f"aviso_{chave}", None)
    if mensagem:
        st.success(mensagem)


# ----------------------------------------------------------------------------
# Leitura de planilhas
# ----------------------------------------------------------------------------
@st.cache_data(show_spinner=False, max_entries=16)
def listar_abas(conteudo: bytes) -> list[str]:
    wb = load_workbook(io.BytesIO(conteudo), read_only=True)
    abas = wb.sheetnames
    wb.close()
    return abas


@st.cache_data(show_spinner="Lendo a planilha...", max_entries=16)
def ler_aba(conteudo: bytes, aba: str) -> pd.DataFrame:
    """Aba inteira, sem cabeçalho. Índice 0 = linha 1 do Excel.

    Lemos com o openpyxl (e não com pd.read_excel) para que a grade seja exatamente
    a mesma que o core.preencher vai escrever, linha por linha.
    """
    wb = load_workbook(io.BytesIO(conteudo), data_only=True)
    linhas = list(wb[aba].iter_rows(values_only=True))
    wb.close()
    return pd.DataFrame(linhas, dtype=object)


# ----------------------------------------------------------------------------
# Detecção automática de cabeçalho e colunas
# ----------------------------------------------------------------------------
@dataclass(frozen=True)
class Campo:
    nome: str
    rotulo: str
    chaves: tuple[tuple[str, ...], ...]   # tentativas em ordem de preferência
    evitar: tuple[str, ...] = ()
    obrigatorio: bool = True
    ajuda: str | None = None


DESC = Campo("desc", "Descrição do item", (("DESCRI", "ESPECIFICA", "DISCRIMINA"),))
VALOR = Campo("valor", "Valor unitário", (("UNIT",), ("VALOR", "PRECO", "CUSTO")), evitar=("TOTAL", "QUANT"))
CAMPOS_LICITACAO = [
    DESC,
    Campo("valor", "Valor unitário (onde o preço será escrito)", VALOR.chaves, VALOR.evitar),
    Campo("qtd", "Quantidade", (("QUANT", "QTD"),), obrigatorio=False,
          ajuda="Linhas sem quantidade são tratadas como títulos de grupo e não são preenchidas."),
    Campo("fornecedor", "Fornecedor / fonte do preço", (("FORNECEDOR", "FONTE", "ORIGEM"),), obrigatorio=False,
          ajuda=f"Em cada item preenchido o app escreve \"{core.FORNECEDOR_SINAPI}\" ou \"{core.FORNECEDOR_BANCO}\", "
                "conforme a origem do preço."),
    Campo("data", "Data do preço", (("DATA",),), obrigatorio=False,
          ajuda="SINAPI: mês de referência da tabela. Banco da empresa: data da cotação."),
]
CAMPOS_SINAPI = [DESC, Campo("valor", "Preço", (("PRECO", "CUSTO", "VALOR"),), evitar=("TOTAL",))]
CAMPOS_CADASTRO = [
    DESC,
    VALOR,
    Campo("codigo", "Código", (("CODIGO", "COD."),), obrigatorio=False),
    Campo("fonte_codigo", "Origem do código", (("FONTE", "BANCO", "BASE", "REFERENCIA"),),
          evitar=("DATA", "PRECO", "VALOR"), obrigatorio=False,
          ajuda="Coluna que diz de qual base é o código: SINAPI, SICRO, ORSE..."),
    Campo("fornecedor", "Fornecedor", (("FORNECEDOR", "EMPRESA", "LICITANTE"),), obrigatorio=False),
    Campo("vendedor", "Vendedor", (("VENDEDOR", "REPRESENTANTE", "CONTATO"),), obrigatorio=False),
    Campo("data", "Data da cotação", (("DATA",),), obrigatorio=False),
]
OPCOES_FONTE_CODIGO = [f"Banco da empresa ({core.FONTE_CODIGO_EMPRESA})", "SINAPI", "Outra base", "Não informar"]
VALOR_FONTE_CODIGO = {OPCOES_FONTE_CODIGO[0]: core.FONTE_CODIGO_EMPRESA, "SINAPI": "SINAPI", "Não informar": None}
GRUPOS_CABECALHO = (DESC.chaves[0], ("UNIT", "VALOR", "PRECO", "CUSTO"), ("QUANT", "QTD"))


def _linha_norm(raw: pd.DataFrame, i: int) -> list[str]:
    return [core.normalizar(v) for v in raw.iloc[i].tolist()]


def detectar_cabecalho(raw: pd.DataFrame, limite: int = 60) -> int:
    """Linha (base 1) que mais parece cabeçalho: precisa ter 'descrição' e, de preferência, valor e quantidade."""
    melhor, melhor_pontos = 1, 0
    for i in range(min(len(raw), limite)):
        celulas = _linha_norm(raw, i)
        grupos = [any(any(k in c for k in grupo) for c in celulas) for grupo in GRUPOS_CABECALHO]
        pontos = sum(grupos) if grupos[0] else 0
        if pontos > melhor_pontos:
            melhor, melhor_pontos = i + 1, pontos
    return melhor


def detectar_coluna(raw: pd.DataFrame, linha_cab: int, campo: Campo) -> int | None:
    if not 1 <= linha_cab <= len(raw):
        return None
    celulas = _linha_norm(raw, linha_cab - 1)
    for tentativa in campo.chaves:
        for j, c in enumerate(celulas):
            if any(k in c for k in tentativa) and not any(e in c for e in campo.evitar):
                return j
    return None


@dataclass
class Selecao:
    raw: pd.DataFrame
    aba: str
    linha_cab: int
    cols: dict[str, int | None]
    chave: str  # prefixo único para widgets que dependem desta seleção


def seletor_planilha(prefixo: str, arquivo, campos: list[Campo], previa: bool = True) -> Selecao | None:
    """Aba, linha do cabeçalho e colunas, já sugeridos automaticamente. O usuário só confere."""
    conteudo = arquivo.getvalue()
    try:
        abas = listar_abas(conteudo)
    except Exception:
        st.error("Não foi possível abrir este arquivo. Abra no Excel, salve como .xlsx e envie de novo.")
        return None

    base = f"{prefixo}_{arquivo.name}_{arquivo.size}"
    c_aba, c_cab = st.columns([3, 1])
    aba = c_aba.selectbox("Aba", abas, key=f"{base}_aba")
    raw = ler_aba(conteudo, aba)
    if raw.empty:
        st.warning("Esta aba está vazia. Escolha outra aba.")
        return None
    linha_cab = int(c_cab.number_input(
        "Linha do cabeçalho", min_value=1, max_value=len(raw), value=detectar_cabecalho(raw),
        key=f"{base}_{aba}_cab", help="Linha onde estão os títulos das colunas (Descrição, Valor...).",
    ))

    chave = f"{base}_{aba}_{linha_cab}"
    titulos = raw.iloc[linha_cab - 1].tolist()
    nomes = {j: f"Coluna {letra(j)}: {core.texto(titulos[j])[:45] or '(sem título)'}" for j in range(raw.shape[1])}
    opcoes = list(range(raw.shape[1]))

    cols: dict[str, int | None] = {}
    for inicio in range(0, len(campos), 3):
        grade = st.columns(3)
        for pos, campo in enumerate(campos[inicio:inicio + 3]):
            sugerida = detectar_coluna(raw, linha_cab, campo)
            if campo.obrigatorio:
                ops, idx = opcoes, (sugerida if sugerida is not None else 0)
            else:
                ops, idx = [None] + opcoes, (0 if sugerida is None else sugerida + 1)
            cols[campo.nome] = grade[pos].selectbox(
                campo.rotulo, ops, index=idx, key=f"{chave}_{campo.nome}", help=campo.ajuda,
                format_func=lambda j: "Não tem" if j is None else nomes[j],
            )

    usadas = [c for c in campos if cols[c.nome] is not None]
    repetidas = {cols[c.nome] for c in usadas if sum(cols[o.nome] == cols[c.nome] for o in usadas) > 1}
    if repetidas:
        st.error(f"A coluna {', '.join(letra(j) for j in sorted(repetidas))} foi escolhida para mais de um campo. Cada campo precisa de uma coluna diferente.")
        return None

    if previa:
        fatia = raw.iloc[linha_cab:linha_cab + 6]
        tabela = pd.DataFrame({c.rotulo: fatia[cols[c.nome]].map(core.texto) for c in usadas})
        tabela.index = [f"Linha {i + 1}" for i in fatia.index]
        with st.expander("Conferir as primeiras linhas lidas"):
            st.dataframe(tabela)

    return Selecao(raw=raw, aba=aba, linha_cab=linha_cab, cols=cols, chave=chave)


# ----------------------------------------------------------------------------
# Banco de dados
# ----------------------------------------------------------------------------
def _url_banco() -> str | None:
    try:
        url = st.secrets.get("DATABASE_URL")
    except Exception:  # sem arquivo de secrets (rodando localmente)
        url = None
    return url or os.environ.get("DATABASE_URL")


@st.cache_resource(show_spinner="Conectando ao banco de preços...")
def obter_engine(versao_esquema: int):
    """A versão faz parte da chave do cache: mudou a estrutura do banco, abre conexão nova e atualiza as tabelas."""
    engine = storage.criar_engine(_url_banco())
    storage.inicializar(engine)
    return engine


@st.cache_data(ttl=300, show_spinner=False)
def carregar_banco(_engine) -> pd.DataFrame:
    return storage.ler_precos(_engine)


# ----------------------------------------------------------------------------
# ----------------------------------------------------------------------------
# Componentes visuais
# ----------------------------------------------------------------------------
def etapa(prefixo: str, numero: int, titulo: str, descricao: str = ""):
    """Bloco de um passo do fluxo, com número em círculo ligado ao próximo passo."""
    caixa = st.container(key=f"etapa_{prefixo}_{numero}")
    desc = f'<div class="atl-etapa-desc">{html.escape(descricao)}</div>' if descricao else ""
    caixa.markdown(
        f'<div class="atl-etapa-cab"><span class="atl-num">{numero}</span>'
        f'<div><div class="atl-etapa-titulo">{html.escape(titulo)}</div>{desc}</div></div>',
        unsafe_allow_html=True,
    )
    return caixa


def estado_vazio(titulo: str, texto_html: str) -> None:
    st.markdown(f'<div class="atl-vazio"><div class="titulo">{html.escape(titulo)}</div><p>{texto_html}</p></div>',
                unsafe_allow_html=True)


def topo(engine, banco: pd.DataFrame | None, erro: Exception | None) -> None:
    if engine is None or banco is None:
        estado, cor = "Banco de preços offline", "#F5B841"
    elif not storage.eh_remoto(engine):
        estado, cor = f"Banco temporário (teste): {fmt_int(len(banco))} preços", "#F5B841"
    else:
        estado, cor = f"Banco de preços: {fmt_int(len(banco))} registros", "#4CC38A"
    dica = f"Erro de conexão: {type(erro).__name__}" if erro is not None else estado
    st.markdown(
        f"""<div class="atl-hero">
          <div class="atl-logo">{logo_html()}</div>
          <div class="atl-hero-texto"><div class="nome">Orçamento de licitações</div>
            <div class="sub">Preencha planilhas com o banco de preços da empresa e a tabela SINAPI</div></div>
          <div class="atl-banco" style="--cor-ponto:{cor}" title="{html.escape(dica)}">{html.escape(estado)}</div>
        </div>""",
        unsafe_allow_html=True,
    )


def barra_cobertura(contagem: pd.Series, total: int) -> None:
    ok = int(contagem.get(core.ST_OK, 0))
    segmentos, legenda = [], []
    for status in ORDEM_STATUS:
        n = int(contagem.get(status, 0))
        if not n:
            continue
        segmentos.append(f'<span style="width:{n / total * 100:.3f}%;background:{COR_BARRA[status]}" '
                         f'title="{n} {LEGENDA_STATUS[status]}"></span>')
        legenda.append(f'<span><i style="background:{COR_BARRA[status]}"></i>{fmt_int(n)} {LEGENDA_STATUS[status]}</span>')
    st.markdown(
        f"""<div class="atl-resultado">
          <p class="atl-resumo"><strong>{fmt_int(ok)}</strong>de {fmt_int(total)} itens preenchidos</p>
          <div class="atl-barra" role="img" aria-label="{ok} de {total} itens preenchidos">{''.join(segmentos)}</div>
          <div class="atl-legenda">{''.join(legenda)}</div>
        </div>""",
        unsafe_allow_html=True,
    )


def _pintar(cores: dict[str, str]):
    def estilo(valor):
        cor = cores.get(valor)
        return f"background-color: {cor}; color: #1A1A1A" if cor else ""
    return estilo


def tabela_relatorio(df: pd.DataFrame) -> None:
    vis = df.copy()
    vis["Valor unitário"] = vis["Valor unitário"].map(brl)
    st.dataframe(
        vis.style.map(_pintar(COR_CELULA_STATUS), subset=["Status"]),
        hide_index=True,
        column_config={
            "Linha Excel": st.column_config.NumberColumn("Linha", format="%d", width="small"),
            "Descrição": st.column_config.TextColumn(width="large"),
        },
    )


PERIODOS = {
    "Todo o período": None,
    "Últimos 30 dias": {"days": 30},
    "Últimos 90 dias": {"days": 90},
    "Últimos 6 meses": {"months": 6},
    "Último ano": {"months": 12},
    "Últimos 2 anos": {"months": 24},
    "Personalizado": "personalizado",
}


def hoje_brasilia() -> date:
    return datetime.now(storage.BRT).date()


def filtros_banco(banco: pd.DataFrame) -> tuple[pd.DataFrame, str]:
    """Controles de período e fornecedor. Devolve o banco filtrado e uma descrição do filtro para o relatório."""
    hoje = hoje_brasilia()
    periodo = st.selectbox("Período das cotações", list(PERIODOS), key="periodo_banco",
                           help="Só os preços cotados dentro deste período são usados no preenchimento.")
    regra = PERIODOS[periodo]
    inicio = fim = None
    if regra == "personalizado":
        intervalo = st.date_input("Intervalo", value=(hoje - timedelta(days=90), hoje), max_value=hoje,
                                  format="DD/MM/YYYY", key="periodo_intervalo")
        if isinstance(intervalo, (tuple, list)) and len(intervalo) == 2:
            inicio, fim = intervalo
        elif isinstance(intervalo, (tuple, list)) and len(intervalo) == 1:
            inicio = intervalo[0]
            st.caption("Escolha também a data final no calendário.")
    elif regra:
        inicio = (pd.Timestamp(hoje) - pd.DateOffset(**regra)).date()

    incluir_sem_data = False
    if inicio is not None or fim is not None:
        sem_data = int(pd.to_datetime(banco["data_cotacao"], errors="coerce").isna().sum())
        if sem_data:
            incluir_sem_data = st.checkbox(f"Incluir {fmt_int(sem_data)} preços cadastrados sem data",
                                           key="incluir_sem_data",
                                           help="Sem data não dá para saber se a cotação está dentro do período.")

    nomes = banco["fornecedor"].fillna("").astype(str).str.strip().replace("", core.SEM_FORNECEDOR)
    with st.expander("Filtrar por fornecedor"):
        fornecedores = st.multiselect("Usar só preços destes fornecedores", sorted(nomes.unique()),
                                      placeholder="Todos os fornecedores", key="filtro_fornecedores")

    filtrado = core.filtrar_banco(banco, inicio, fim, incluir_sem_data, fornecedores)

    partes = []
    if inicio is not None or fim is not None:
        de = f"{inicio:%d/%m/%Y}" if inicio else "início"
        ate = f"{(fim or hoje):%d/%m/%Y}"
        partes.append(f"cotações de {de} a {ate}" + (" (incluindo sem data)" if incluir_sem_data else ""))
    if fornecedores:
        partes.append(f"{len(fornecedores)} fornecedor(es)")
    descricao = ", ".join(partes) if partes else "todo o período"

    if filtrado.empty:
        st.warning("Nenhum preço do banco dentro deste filtro. O preenchimento vai usar só o SINAPI.")
    else:
        st.caption(f"{fmt_int(len(filtrado))} de {fmt_int(len(banco))} preços dentro do filtro ({descricao}).")
    return filtrado, descricao


COMO_FUNCIONA = """<div class="atl-como">
  <div><b>Envie a planilha</b><span>O arquivo .xlsx da licitação, do jeito que veio do órgão.</span></div>
  <div><b>Confira as colunas</b><span>Descrição, valor e quantidade são encontrados sozinhos. Você só confirma.</span></div>
  <div><b>Baixe preenchida</b><span>Os itens sem preço saem destacados em amarelo ou laranja para revisão.</span></div>
</div>"""

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ----------------------------------------------------------------------------
# Tela 1: preencher planilha
# ----------------------------------------------------------------------------
def tela_preencher(banco: pd.DataFrame | None) -> None:
    with etapa("lic", 1, "Planilha da licitação", "Envie a planilha que será preenchida e confira as colunas encontradas."):
        arquivo = st.file_uploader(
            "Planilha da licitação", type=["xlsx", "xlsm"], key="up_licitacao", label_visibility="collapsed",
            help="Arquivos .xls antigos: abra no Excel e salve como .xlsx.",
        )
        if arquivo is None:
            st.markdown(COMO_FUNCIONA, unsafe_allow_html=True)
            return
        selecao = seletor_planilha("lic", arquivo, CAMPOS_LICITACAO)
    if selecao is None:
        return

    tem_banco = banco is not None and not banco.empty
    banco_uso, descricao_filtro = banco, ""
    estrategia = core.ESTRATEGIAS[0]
    cfg_sinapi, mes_sinapi = None, None
    with etapa("lic", 2, "Fontes de preço", "O banco da empresa é consultado primeiro. O SINAPI cobre o que faltar."):
        col_banco, col_sinapi = st.columns(2, gap="large")
        with col_banco:
            st.markdown("**Banco de preços da empresa**")
            if banco is None:
                st.caption("Offline no momento. O preenchimento vai usar só o SINAPI.")
            elif banco.empty:
                st.caption("Ainda sem preços cadastrados. Use a aba Cadastrar preços para alimentar o banco.")
            else:
                banco_uso, descricao_filtro = filtros_banco(banco)
                tem_banco = not banco_uso.empty
                estrategia = st.radio(
                    "Quando o item tiver mais de um preço, usar", core.ESTRATEGIAS, horizontal=True, key="estrategia",
                    help="Mais recente: a cotação de data mais nova. Menor preço: o mais barato. Mediana: o valor do meio.",
                )
        with col_sinapi:
            st.markdown("**Tabela SINAPI** (opcional)")
            arquivo_sinapi = st.file_uploader("Planilha do SINAPI", type=["xlsx", "xlsm"], key="up_sinapi",
                                              label_visibility="collapsed")
        if arquivo_sinapi is not None:
            with st.expander("Colunas do SINAPI (encontradas automaticamente)"):
                sel_sinapi = seletor_planilha("sin", arquivo_sinapi, CAMPOS_SINAPI, previa=False)
            if sel_sinapi is not None:
                cfg_sinapi = core.ConfigPlanilha(
                    raw=sel_sinapi.raw, aba=sel_sinapi.aba, linha_cab=sel_sinapi.linha_cab,
                    col_desc=sel_sinapi.cols["desc"], col_valor=sel_sinapi.cols["valor"],
                )
                c_mes, c_info = st.columns([1, 2])
                sugestao = core.detectar_mes_sinapi(arquivo_sinapi.name, sel_sinapi.raw)
                mes_texto = c_mes.text_input(
                    "Mês de referência do SINAPI", value=f"{sugestao:%m/%Y}" if sugestao else "",
                    placeholder="MM/AAAA", key=f"{sel_sinapi.chave}_mes",
                    help="Vai para a coluna de data dos itens preenchidos pelo SINAPI.",
                )
                mes_sinapi = core.para_mes(mes_texto)
                if mes_texto.strip() and mes_sinapi is None:
                    c_info.warning("Mês não reconhecido. Use o formato MM/AAAA, por exemplo 09/2026.")
                else:
                    c_info.caption(
                        f"Descrição na coluna {letra(cfg_sinapi.col_desc)}, preço na coluna {letra(cfg_sinapi.col_valor)}, "
                        f"cabeçalho na linha {cfg_sinapi.linha_cab}."
                        + ("" if sugestao else " Não encontrei o mês na tabela: informe ao lado.")
                    )

    assinatura = f"{arquivo.name}-{arquivo.size}"
    with etapa("lic", 3, "Preencher", "Gera uma cópia preenchida. A planilha original não é alterada."):
        sobrescrever = st.toggle(
            "Substituir valores que já estão na planilha", value=False,
            help="Desligado: itens que já têm valor ficam como estão e aparecem no relatório como mantidos.",
        )
        sem_fonte = not tem_banco and cfg_sinapi is None
        if sem_fonte:
            st.caption("Para preencher, envie a tabela SINAPI no passo 2 ou cadastre preços no banco.")
        elif selecao.cols["fornecedor"] is None and selecao.cols["data"] is None:
            st.caption("Para registrar de onde veio cada preço, escolha as colunas de fornecedor e data no passo 1.")
        if st.button("Preencher planilha", type="primary", disabled=sem_fonte, icon=":material/edit_note:"):
            cfg = core.ConfigPlanilha(
                raw=selecao.raw, aba=selecao.aba, linha_cab=selecao.linha_cab,
                col_desc=selecao.cols["desc"], col_valor=selecao.cols["valor"], col_qtd=selecao.cols["qtd"],
                col_fornecedor=selecao.cols["fornecedor"], col_data=selecao.cols["data"],
            )
            with st.spinner("Procurando os preços..."):
                interno = core.montar_indice_banco(banco_uso, estrategia) if tem_banco else core.Indice()
                sinapi = core.montar_indice(cfg_sinapi) if cfg_sinapi is not None else core.Indice()
                try:
                    saida, relatorio = core.preencher(arquivo.getvalue(), cfg, interno, sinapi, sobrescrever,
                                                      mes_sinapi=mes_sinapi)
                except ValueError as erro:
                    st.error(f"{erro} Confira a aba e a linha do cabeçalho e tente de novo.")
                    return
            st.session_state["resultado"] = {
                "assinatura": assinatura, "planilha": saida, "relatorio": relatorio,
                "filtro": f"Banco da empresa: {descricao_filtro}, critério {estrategia.lower()}." if tem_banco else "",
                "nome": f"{Path(arquivo.name).stem}_preenchida.xlsx", "xlsm": arquivo.name.lower().endswith(".xlsm"),
            }

    resultado = st.session_state.get("resultado")
    if resultado and resultado["assinatura"] == assinatura:
        with etapa("lic", 4, "Resultado", "Baixe a planilha e revise os itens destacados."):
            mostrar_resultado(resultado)


def mostrar_resultado(resultado: dict) -> None:
    relatorio: pd.DataFrame = resultado["relatorio"]
    if relatorio.empty:
        st.warning("Nenhum item foi encontrado abaixo do cabeçalho. Confira a linha do cabeçalho, a coluna de "
                   "descrição e, se escolheu, a coluna de quantidade.")
        return

    barra_cobertura(relatorio["Status"].value_counts(), len(relatorio))
    if resultado.get("filtro"):
        st.caption(resultado["filtro"])
    c1, c2, _ = st.columns([1.4, 1.1, 1.5])
    c1.download_button("Baixar planilha preenchida", resultado["planilha"], file_name=resultado["nome"],
                       type="primary", mime=XLSX_MIME, icon=":material/download:")
    c2.download_button("Baixar relatório", core.df_para_xlsx(relatorio), mime=XLSX_MIME, icon=":material/description:",
                       file_name=resultado["nome"].replace("_preenchida", "_relatorio"))
    if resultado["xlsm"]:
        st.caption("A planilha original tinha macros. A versão preenchida sai em .xlsx, sem as macros.")

    pendentes = relatorio[relatorio["Status"].isin([core.ST_NAO, core.ST_AMB, core.ST_MESCLADA])]
    opcoes = [f"Itens para revisar ({fmt_int(len(pendentes))})", f"Todos os itens ({fmt_int(len(relatorio))})"]
    escolha = st.radio("Mostrar", opcoes, horizontal=True, key="filtro_resultado", label_visibility="collapsed")
    if escolha == opcoes[0]:
        if pendentes.empty:
            st.success("Todos os itens receberam preço. Nada para revisar.")
            return
        st.caption("Estes itens ficaram destacados com a mesma cor na planilha baixada e precisam de preço manual.")
        tabela_relatorio(pendentes)
    else:
        tabela_relatorio(relatorio)


# ----------------------------------------------------------------------------
# Tela 2: cadastrar preços
# ----------------------------------------------------------------------------
BANCO_OFFLINE = ("O cadastro e a consulta voltam a funcionar quando o banco for configurado. Enquanto isso, "
                 "a aba <b>Preencher planilha</b> funciona normalmente com a tabela SINAPI.")


def tela_cadastrar(engine) -> None:
    mostrar_aviso("cadastro")
    if engine is None:
        estado_vazio("Banco de preços offline", BANCO_OFFLINE)
        return
    if not storage.eh_remoto(engine):
        st.caption("Modo de teste: este banco é temporário e é apagado quando o app reinicia.")

    with etapa("cad", 1, "Planilha de uma licitação passada",
               "Cada linha com descrição e valor unitário vira um preço no banco."):
        arquivo = st.file_uploader("Planilha com preços cotados ou contratados", type=["xlsx", "xlsm"],
                                   key="up_cadastro", label_visibility="collapsed")
        if arquivo is None:
            st.caption("Antes de gravar, você confere o que é novo e o que já está cadastrado.")
            return
        selecao = seletor_planilha("cad", arquivo, CAMPOS_CADASTRO)
    if selecao is None:
        return
    cols = selecao.cols

    with etapa("cad", 2, "Dados da licitação", "Informações que valem para todos os preços desta planilha."):
        c1, c2 = st.columns(2)
        origem = c1.text_input(
            "Identificação da licitação", value=Path(arquivo.name).stem, key=f"{selecao.chave}_fixo_origem",
            help="Ex.: Pregão 12/2025, TRE-DF. Fica gravado junto com cada preço.",
        )
        data_fixa = None
        if cols["data"] is None:
            data_fixa = c2.date_input("Data da cotação", value=None, format="DD/MM/YYYY",
                                      key=f"{selecao.chave}_fixo_data")
        else:
            c2.caption(f"Data da cotação lida da coluna {letra(cols['data'])}.")
        c3, c4 = st.columns(2)
        fornecedor_fixo = vendedor_fixo = None
        if cols["fornecedor"] is None:
            fornecedor_fixo = c3.text_input("Fornecedor", key=f"{selecao.chave}_fixo_fornecedor")
        else:
            c3.caption(f"Fornecedor lido da coluna {letra(cols['fornecedor'])}.")
        if cols["vendedor"] is None:
            vendedor_fixo = c4.text_input("Vendedor", key=f"{selecao.chave}_fixo_vendedor")
        else:
            c4.caption(f"Vendedor lido da coluna {letra(cols['vendedor'])}.")

        fonte_codigo_fixa = None
        if cols["codigo"] is None:
            if cols["fonte_codigo"] is not None:
                st.caption("A origem do código só é gravada quando a coluna de código também é escolhida no passo 1.")
        elif cols["fonte_codigo"] is not None:
            st.caption(f"Origem do código lida da coluna {letra(cols['fonte_codigo'])}.")
        else:
            c5, c6 = st.columns(2)
            escolha_fonte = c5.radio(
                "Origem do código (vale para todos os itens)", OPCOES_FONTE_CODIGO, index=len(OPCOES_FONTE_CODIGO) - 1,
                key=f"{selecao.chave}_fixo_fonte_codigo",
                help="A planilha não tem coluna com a base do código. Escolha de onde vêm os códigos desta planilha.",
            )
            if escolha_fonte == "Outra base":
                fonte_codigo_fixa = c6.text_input("Nome da base", placeholder="Ex.: SICRO, ORSE",
                                                  key=f"{selecao.chave}_fixo_fonte_outra")
            else:
                fonte_codigo_fixa = VALOR_FONTE_CODIGO[escolha_fonte]

    mapa = core.MapaIngestao(
        linha_cab=selecao.linha_cab, col_desc=cols["desc"], col_valor=cols["valor"], col_codigo=cols["codigo"],
        col_fornecedor=cols["fornecedor"], fornecedor_fixo=(fornecedor_fixo or "").strip() or None,
        col_vendedor=cols["vendedor"], vendedor_fixo=(vendedor_fixo or "").strip() or None,
        col_data=cols["data"], data_fixa=data_fixa,
        col_fonte_codigo=cols["fonte_codigo"] if cols["codigo"] is not None else None,
        fonte_codigo_fixa=fonte_codigo_fixa,
    )
    validos, ignorados = core.extrair_registros(selecao.raw, mapa)

    with etapa("cad", 3, "Conferir e gravar", "Só os preços novos são gravados. Repetidos são ignorados."):
        if not validos:
            st.warning("Nenhuma linha com descrição e valor unitário válido. Confira a linha do cabeçalho e as colunas.")
            if ignorados:
                st.dataframe(pd.DataFrame(ignorados), hide_index=True)
            return
        try:
            situacoes = storage.classificar(engine, validos)
        except Exception as erro:
            st.error(f"Não foi possível consultar o banco: {erro}")
            return

        tabela = pd.DataFrame(validos)
        tabela.insert(0, "Situação", situacoes)
        novos = int((tabela["Situação"] == storage.SIT_NOVO).sum())
        existentes = int((tabela["Situação"] == storage.SIT_EXISTE).sum())
        repetidos = int((tabela["Situação"] == storage.SIT_REPETIDO).sum())
        resumo = (f"**{fmt_int(len(tabela))} preços encontrados:** {fmt_int(novos)} novos, "
                  f"{fmt_int(existentes)} já estão no banco e {fmt_int(repetidos)} se repetem na própria planilha.")
        if ignorados:
            resumo += f" {fmt_int(len(ignorados))} linhas foram ignoradas."
        st.markdown(resumo)

        vis = pd.DataFrame({
            "Situação": tabela["Situação"],
            "Linha": tabela["linha_excel"],
            "Descrição": tabela["descricao"],
            "Valor unitário": tabela["valor_unitario"].map(brl),
            "Fornecedor": tabela["fornecedor"],
            "Data": pd.to_datetime(tabela["data_cotacao"], errors="coerce").dt.strftime("%d/%m/%Y").fillna(""),
        })
        if cols["codigo"] is not None:
            vis.insert(2, "Código", tabela["codigo"])
            vis.insert(3, "Origem do código", tabela["fonte_codigo"])
        st.dataframe(vis.style.map(_pintar(COR_SITUACAO), subset=["Situação"]), hide_index=True,
                     column_config={"Descrição": st.column_config.TextColumn(width="large")})
        if ignorados:
            with st.expander(f"Linhas ignoradas ({fmt_int(len(ignorados))})"):
                st.dataframe(pd.DataFrame(ignorados), hide_index=True)

        if novos == 0:
            st.caption("Todos os preços desta planilha já estão no banco.")
        if st.button(f"Gravar {fmt_int(novos)} preços novos", type="primary", disabled=novos == 0,
                     icon=":material/save:"):
            if not origem.strip():
                st.error("Informe a identificação da licitação antes de gravar.")
                return
            with st.spinner("Gravando..."):
                resultado = storage.inserir(engine, validos, origem.strip())
            carregar_banco.clear()
            avisar("cadastro", f"{fmt_int(resultado['inseridos'])} preços gravados de \"{origem.strip()}\". "
                               "Se precisar desfazer, remova esta importação na aba Consultar banco.")
            st.rerun()


# ----------------------------------------------------------------------------
# Tela 3: consultar banco
# ----------------------------------------------------------------------------
def tela_consultar(engine, banco: pd.DataFrame | None) -> None:
    mostrar_aviso("consulta")
    if engine is None or banco is None:
        estado_vazio("Banco de preços offline", BANCO_OFFLINE)
        return
    if banco.empty:
        estado_vazio("O banco ainda está vazio",
                     "Cadastre a primeira planilha de licitação na aba <b>Cadastrar preços</b>. "
                     "Os preços aparecem aqui para busca e exportação.")
        return

    busca = st.text_input("Buscar item", placeholder="Ex.: cabo flexível 2,5 mm²", icon=":material/search:",
                          help="Mostra os itens que contêm todas as palavras digitadas, sem diferenciar acentos.")
    dados = banco
    termos = core.normalizar(busca).split()
    if termos:
        descricoes = dados["descricao_norm"].fillna("")
        filtro = pd.Series(True, index=dados.index)
        for termo in termos:
            filtro &= descricoes.str.contains(termo, regex=False)
        dados = dados[filtro]

    datas = pd.to_datetime(dados["data_cotacao"], errors="coerce")
    exportar = pd.DataFrame({
        "Descrição": dados["descricao"],
        "Valor unitário": dados["valor_unitario"],
        "Fornecedor": dados["fornecedor"],
        "Vendedor": dados["vendedor"],
        "Data da cotação": datas.dt.date,
        "Licitação": dados["origem"],
        "Código": dados["codigo"],
        "Origem do código": dados["fonte_codigo"] if "fonte_codigo" in dados else None,
    }).assign(_ordem=dados["descricao_norm"], _data=datas)
    exportar = exportar.sort_values(["_ordem", "_data"], ascending=[True, False]).drop(columns=["_ordem", "_data"])

    c_info, c_botao = st.columns([3, 1])
    c_info.caption(f"{fmt_int(len(exportar))} de {fmt_int(len(banco))} registros")
    c_botao.download_button("Exportar (.xlsx)", core.df_para_xlsx(exportar, "Banco de preços"),
                            file_name="banco_de_precos.xlsx", mime=XLSX_MIME, icon=":material/download:")
    vis = exportar.assign(**{
        "Valor unitário": exportar["Valor unitário"].map(brl),
        "Data da cotação": pd.to_datetime(exportar["Data da cotação"]).dt.strftime("%d/%m/%Y").fillna(""),
    })
    st.dataframe(vis, hide_index=True, column_config={"Descrição": st.column_config.TextColumn(width="large")})

    with st.expander("Importações (desfazer um cadastro)"):
        st.caption("Cada planilha cadastrada é uma importação. Se algo foi gravado errado, remova a importação "
                   "inteira e cadastre a planilha de novo.")
        lotes = storage.listar_lotes(engine)
        lotes["importado_em"] = pd.to_datetime(lotes["importado_em"], errors="coerce")
        st.dataframe(
            pd.DataFrame({
                "Licitação": lotes["origem"],
                "Importada em": lotes["importado_em"].dt.strftime("%d/%m/%Y %H:%M"),
                "Preços": lotes["registros"],
            }),
            hide_index=True,
        )
        info = {r.lote: r for r in lotes.itertuples()}
        lote = st.selectbox(
            "Importação para remover", list(info), index=None, placeholder="Escolha uma importação",
            format_func=lambda l: f"{info[l].origem} ({fmt_int(info[l].registros)} preços, "
                                  f"{info[l].importado_em:%d/%m/%Y %H:%M})",
        )
        if lote is not None:
            confirmar = st.checkbox(f"Confirmo que quero apagar os {fmt_int(info[lote].registros)} preços desta importação")
            if st.button("Remover importação", disabled=not confirmar, icon=":material/delete:"):
                removidos = storage.remover_lote(engine, lote)
                carregar_banco.clear()
                avisar("consulta", f"Importação removida: {fmt_int(removidos)} preços apagados do banco.")
                st.rerun()


# ----------------------------------------------------------------------------
# Página
# ----------------------------------------------------------------------------
def main() -> None:
    logo = arquivo_logo()
    st.set_page_config(
        page_title="Orçamento de licitações | Atlântico Engenharia",
        page_icon=str(logo) if logo and logo.suffix != ".svg" else "📋",
        layout="wide",
    )
    st.markdown(CSS, unsafe_allow_html=True)

    engine, banco, erro = None, None, None
    try:
        engine = obter_engine(storage.VERSAO_ESQUEMA)
        banco = carregar_banco(engine)
    except Exception as e:  # o app continua funcionando só com o SINAPI
        engine, erro = None, e

    topo(engine, banco, erro)
    aba_preencher, aba_cadastrar, aba_consultar = st.tabs([
        ":material/edit_document: Preencher planilha",
        ":material/library_add: Cadastrar preços",
        ":material/search: Consultar banco",
    ])
    with aba_preencher:
        tela_preencher(banco)
    with aba_cadastrar:
        tela_cadastrar(engine)
    with aba_consultar:
        tela_consultar(engine, banco)

    st.markdown('<div class="atl-rodape">Ferramenta interna de orçamento da Atlântico Engenharia e Inovação.</div>',
                unsafe_allow_html=True)


if __name__ == "__main__":
    main()
