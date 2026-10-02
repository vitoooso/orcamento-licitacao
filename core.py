"""
core.py — regras de negócio (sem Streamlit e sem banco de dados).

Normalização, índices de preço, extração de planilhas de licitações passadas
e preenchimento da planilha da licitação.
"""
from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from numbers import Number

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill

# ----------------------------------------------------------------------------
# Constantes
# ----------------------------------------------------------------------------
TOLERANCIA_VALOR = 0.005  # preços que diferem menos que meio centavo = mesmo preço
ESTRATEGIAS = ("Mais recente", "Menor preço", "Mediana")

FILL_AMARELO = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
FILL_LARANJA = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")

ST_OK = "Preenchido"
ST_NAO = "Não encontrado"
ST_AMB = "Ambíguo (valores conflitantes)"
ST_JA = "Já preenchido (mantido)"
ST_MESCLADA = "Célula de valor mesclada"
FONTE_INTERNO = "Banco interno"
FONTE_SINAPI = "SINAPI"
COLUNAS_RELATORIO = ["Linha Excel", "Descrição", "Status", "Fonte", "Valor unitário", "Observação"]
COR_STATUS = {ST_NAO: "#FFFF00", ST_AMB: "#FFC000", ST_MESCLADA: "#FFC000", ST_JA: "#D9D9D9"}


@dataclass
class ConfigPlanilha:
    raw: pd.DataFrame        # aba inteira, sem cabeçalho (índice 0 = linha 1 do Excel)
    aba: str
    linha_cab: int           # nº da linha do cabeçalho no Excel (base 1)
    col_desc: int            # índice da coluna de descrição (base 0)
    col_valor: int           # índice da coluna de valor (base 0)
    col_qtd: int | None = None


@dataclass
class Indice:
    """Resultado de uma base de preços já pronta para consulta."""
    valores: dict[str, float] = field(default_factory=dict)          # chave normalizada -> preço
    conflitos: dict[str, list[float]] = field(default_factory=dict)  # chave -> preços divergentes
    notas: dict[str, str] = field(default_factory=dict)              # chave -> texto de auditoria


@dataclass
class MapaIngestao:
    """Como ler uma planilha de licitação passada para alimentar o banco."""
    linha_cab: int
    col_desc: int
    col_valor: int
    col_codigo: int | None = None
    col_fornecedor: int | None = None
    fornecedor_fixo: str | None = None
    col_vendedor: int | None = None
    vendedor_fixo: str | None = None
    col_data: int | None = None
    data_fixa: date | None = None


# ----------------------------------------------------------------------------
# Normalização e conversões
# ----------------------------------------------------------------------------
def normalizar(texto) -> str:
    """Remove acentos, converte para MAIÚSCULAS e limpa espaços.

    NFD separa a letra do acento ("é" -> "e" + "´") e removemos só as marcas
    combinantes (categoria Mn). Usamos NFD (e não NFKD) de propósito: NFKD também
    converteria "m²" em "m2" e "º" em "o", unificando especificações que podem
    ser diferentes. Espaços/tabs/quebras de linha/NBSP viram um único espaço.
    """
    if texto is None or pd.isna(texto):
        return ""
    s = unicodedata.normalize("NFD", str(texto))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s.upper()).strip()


def texto(valor) -> str:
    """Texto limpo de uma célula ('' para vazio; 1234.0 vira '1234')."""
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def para_numero(valor) -> float | None:
    """Converte para float (aceita 'R$ 1.234,56'). Vazio, NaN, '-', zero e negativos -> None."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, Number):
        numero = float(valor)
    else:
        txt = re.sub(r"[^\d,.\-]", "", str(valor))
        if "," in txt:  # formato brasileiro: '.' = milhar, ',' = decimal
            txt = txt.replace(".", "").replace(",", ".")
        try:
            numero = float(txt)
        except ValueError:
            return None
    return numero if numero == numero and numero > 0 else None  # numero == numero descarta NaN


def para_data(valor) -> date | None:
    """Converte célula em data. Aceita datetime, serial do Excel e textos dd/mm/aaaa ou aaaa-mm-dd."""
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, Number) and not isinstance(valor, bool):
        if 1 <= float(valor) <= 80000:  # serial do Excel
            return (datetime(1899, 12, 30) + timedelta(days=float(valor))).date()
        return None
    partes = str(valor).strip().split()
    if not partes:
        return None
    for formato in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(partes[0], formato).date()
        except ValueError:
            continue
    return None


def _mesma_celula(valor_openpyxl, valor_pandas) -> bool:
    """Trava de segurança: confere se a linha lida pelo pandas é a mesma que o openpyxl vai escrever."""
    if isinstance(valor_openpyxl, str) and valor_openpyxl.startswith("="):
        return True  # fórmula: o openpyxl não expõe o resultado calculado
    if normalizar(valor_openpyxl) == normalizar(valor_pandas):
        return True
    return isinstance(valor_openpyxl, Number) and isinstance(valor_pandas, Number) and valor_openpyxl == valor_pandas


def _fmt_conflito(valores: list[float]) -> str:
    return "Valores conflitantes: " + "; ".join(f"{v:.2f}".replace(".", ",") for v in valores)


# ----------------------------------------------------------------------------
# Índices de preço
# ----------------------------------------------------------------------------
def montar_indice(cfg: ConfigPlanilha) -> Indice:
    """Índice a partir de uma planilha (usado no SINAPI).

    Se a mesma descrição aparece com preços diferentes, NÃO escolhemos um: vai para 'conflitos'.
    """
    dados = cfg.raw.iloc[cfg.linha_cab:, [cfg.col_desc, cfg.col_valor]].copy()
    dados.columns = ["desc", "valor"]
    dados["chave"] = dados["desc"].map(normalizar)
    dados["valor"] = dados["valor"].map(para_numero)
    dados = dados[(dados["chave"] != "") & dados["valor"].notna()]

    idx = Indice()
    for chave, grupo in dados.groupby("chave")["valor"]:
        valores = grupo.tolist()
        if max(valores) - min(valores) <= TOLERANCIA_VALOR:
            idx.valores[chave] = valores[0]
        else:
            idx.conflitos[chave] = sorted(set(valores))
    return idx


def _nota(n: int, estrategia: str, ref) -> str:
    partes = [f"{n} registro(s) no banco · critério: {estrategia.lower()}"]
    if ref is not None:
        if isinstance(ref["fornecedor"], str) and ref["fornecedor"].strip():
            partes.append(f"fornecedor: {ref['fornecedor']}")
        if pd.notna(ref["data_cotacao"]):
            partes.append(f"data: {ref['data_cotacao']:%d/%m/%Y}")
    return " · ".join(partes)


def montar_indice_banco(df: pd.DataFrame, estrategia: str) -> Indice:
    """Índice a partir do banco de preços, que pode ter VÁRIOS preços para a mesma descrição.

    A estratégia decide qual preço usar: mais recente, menor preço ou mediana.
    """
    idx = Indice()
    if df.empty:
        return idx
    d = df[["id", "descricao_norm", "valor_unitario", "fornecedor", "data_cotacao"]].copy()
    d["data_cotacao"] = pd.to_datetime(d["data_cotacao"], errors="coerce")
    for chave, g in d.groupby("descricao_norm"):
        if estrategia == "Menor preço":
            ref = g.loc[g["valor_unitario"].idxmin()]
            valor = float(ref["valor_unitario"])
        elif estrategia == "Mediana":
            ref, valor = None, float(g["valor_unitario"].median())
        else:  # Mais recente: maior data; sem data fica por último; empate = cadastrado por último
            ref = g.sort_values(["data_cotacao", "id"], ascending=False, na_position="last").iloc[0]
            valor = float(ref["valor_unitario"])
        idx.valores[chave] = valor
        idx.notas[chave] = _nota(len(g), estrategia, ref)
    return idx


# ----------------------------------------------------------------------------
# Extração de planilhas de licitações passadas (para alimentar o banco)
# ----------------------------------------------------------------------------
def extrair_registros(raw: pd.DataFrame, m: MapaIngestao) -> tuple[list[dict], list[dict]]:
    """Retorna (válidos, ignorados). Linhas totalmente vazias são descartadas em silêncio."""
    validos: list[dict] = []
    ignorados: list[dict] = []

    def celula(i: int, col: int | None):
        return None if col is None else raw.iat[i, col]

    for i in range(m.linha_cab, len(raw)):
        descricao = texto(raw.iat[i, m.col_desc])
        bruto = raw.iat[i, m.col_valor]
        valor = para_numero(bruto)
        if not descricao and not texto(bruto):
            continue
        base = {"Linha Excel": i + 1, "Descrição": descricao, "Valor (original)": texto(bruto)}
        if not descricao:
            ignorados.append({**base, "Motivo": "Sem descrição"})
            continue
        if valor is None:
            ignorados.append({**base, "Motivo": "Sem valor unitário válido (vazio, zero ou texto)"})
            continue

        fornecedor = texto(celula(i, m.col_fornecedor)) if m.col_fornecedor is not None else (m.fornecedor_fixo or "")
        vendedor = texto(celula(i, m.col_vendedor)) if m.col_vendedor is not None else (m.vendedor_fixo or "")
        data = para_data(celula(i, m.col_data)) if m.col_data is not None else m.data_fixa
        validos.append({
            "linha_excel": i + 1,
            "codigo": texto(celula(i, m.col_codigo)),
            "descricao": descricao,
            "valor_unitario": valor,
            "fornecedor": fornecedor,
            "vendedor": vendedor,
            "data_cotacao": data,
        })
    return validos, ignorados


# ----------------------------------------------------------------------------
# Preenchimento (escrita com openpyxl para preservar a formatação original)
# ----------------------------------------------------------------------------
def preencher(conteudo_lic: bytes, cfg: ConfigPlanilha, interno: Indice, sinapi: Indice, sobrescrever: bool):
    wb = load_workbook(io.BytesIO(conteudo_lic))
    ws = wb[cfg.aba]
    registros = []

    for i in range(cfg.linha_cab, len(cfg.raw)):  # i base 0 -> linha Excel = i + 1
        descricao = cfg.raw.iat[i, cfg.col_desc]
        chave = normalizar(descricao)
        if not chave:
            continue
        if cfg.col_qtd is not None and not normalizar(cfg.raw.iat[i, cfg.col_qtd]):
            continue  # sem quantidade = título de grupo, não é item

        linha = i + 1
        c_desc = ws.cell(row=linha, column=cfg.col_desc + 1)
        c_valor = ws.cell(row=linha, column=cfg.col_valor + 1)
        if not _mesma_celula(c_desc.value, descricao):
            raise ValueError(
                f"Desalinhamento na linha {linha}: o pandas leu '{descricao}', mas o Excel tem "
                f"'{c_desc.value}'. Nada foi gerado para evitar preencher a linha errada."
            )

        # Hierarquia: banco interno -> SINAPI (correspondência exata da chave normalizada)
        valor, fonte, status, obs = None, "", ST_NAO, ""
        if chave in interno.valores:
            valor, fonte, status, obs = interno.valores[chave], FONTE_INTERNO, ST_OK, interno.notas.get(chave, "")
        elif chave in interno.conflitos:  # existe no interno com preços conflitantes: não cai para o SINAPI
            fonte, status, obs = FONTE_INTERNO, ST_AMB, _fmt_conflito(interno.conflitos[chave])
        elif chave in sinapi.valores:
            valor, fonte, status = sinapi.valores[chave], FONTE_SINAPI, ST_OK
        elif chave in sinapi.conflitos:
            fonte, status, obs = FONTE_SINAPI, ST_AMB, _fmt_conflito(sinapi.conflitos[chave])

        if isinstance(c_valor, MergedCell):
            valor, status = None, ST_MESCLADA
        elif status == ST_OK and c_valor.value not in (None, "") and not sobrescrever:
            valor, status, obs = None, ST_JA, f"Valor existente: {c_valor.value}"

        if status == ST_OK:
            c_valor.value = valor
            if c_valor.number_format == "General":
                c_valor.number_format = "#,##0.00"
        elif status in (ST_NAO, ST_AMB, ST_MESCLADA):  # exige preenchimento manual -> destaca
            fill = FILL_AMARELO if status == ST_NAO else FILL_LARANJA
            c_desc.fill = fill
            if not isinstance(c_valor, MergedCell):
                c_valor.fill = fill

        registros.append({
            "Linha Excel": linha, "Descrição": str(descricao), "Status": status,
            "Fonte": fonte, "Valor unitário": valor, "Observação": obs,
        })

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue(), pd.DataFrame(registros, columns=COLUNAS_RELATORIO)


def df_para_xlsx(df: pd.DataFrame, nome_aba: str = "Relatório") -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as escritor:
        df.to_excel(escritor, index=False, sheet_name=nome_aba)
    return buf.getvalue()
